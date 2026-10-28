#!/usr/bin/env python3
"""Build the project GDD Word document from its Markdown source."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT_NAME = "Microsoft YaHei"
CODE_FONT = "Consolas"
ACCENT = "1F5F5B"
DARK = "1E2B2A"
LIGHT_FILL = "EAF2EF"


def set_run_font(run, name: str = FONT_NAME, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction_text, separate, text, end])


def add_inline_runs(paragraph, text: str) -> None:
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_run_font(run)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.color.rgb = RGBColor(68, 68, 68)
            set_run_font(run, CODE_FONT)
        else:
            run = paragraph.add_run(token)
            set_run_font(run)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row_index, row in enumerate(normalized):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_runs(paragraph, value)
            if row_index == 0:
                set_cell_shading(cell, ACCENT)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_FILL)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in (("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 12)):
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(ACCENT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("《逆岁：四契》立项报告")
    set_run_font(header_run, size=9)
    header_run.font.color.rgb = RGBColor(105, 105, 105)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("v1.6  ·  ")
    set_run_font(footer_run, size=9)
    add_field(footer, "PAGE", "1")

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(document: Document, title: str) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(110)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run("《逆岁：四契》")
    set_run_font(title_run, size=30)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(DARK)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("立项报告 / 游戏设计文档")
    set_run_font(subtitle_run, size=18)
    subtitle_run.bold = True
    subtitle_run.font.color.rgb = RGBColor.from_string(ACCENT)

    divider = document.add_paragraph()
    divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    divider_run = divider.add_run("白虎  ·  朱雀  ·  玄武  ·  青龙")
    set_run_font(divider_run, size=12)
    divider_run.font.color.rgb = RGBColor(110, 92, 54)

    info_lines = [
        "Godot 4.7 / 2D 横版剧情动作游戏",
        "四象岁契 · 神兽面具 · 战斗手感重制",
        "版本：v1.6",
        "日期：2026-10-24",
        "状态：主方案，可用于垂直切片制作",
    ]
    for index, line in enumerate(info_lines):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(38 if index == 0 else 4)
        run = paragraph.add_run(line)
        set_run_font(run, size=10.5)

    document.add_page_break()


def add_toc(document: Document) -> None:
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("目录")
    toc = document.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "打开 Word 后更新目录")
    document.add_page_break()


def resolve_image(source_dir: Path, relative_path: str) -> Path | None:
    candidate = (source_dir / relative_path).resolve()
    return candidate if candidate.is_file() else None


def parse_markdown(document: Document, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    source_dir = markdown_path.parent
    i = 0
    first_h1_seen = False

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line or line == "---":
            i += 1
            continue

        if line.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            paragraph = document.add_paragraph(style="No Spacing")
            paragraph.paragraph_format.left_indent = Inches(0.2)
            run = paragraph.add_run("\n".join(code_lines))
            set_run_font(run, CODE_FONT, 9)
            run.font.color.rgb = RGBColor(55, 55, 55)
            i += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            alt_text, relative_path = image_match.groups()
            image_path = resolve_image(source_dir, relative_path)
            if image_path is not None:
                document.add_picture(str(image_path), width=Inches(6.2))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_run = caption.add_run(alt_text)
                set_run_font(caption_run, size=9)
                caption_run.italic = True
            else:
                paragraph = document.add_paragraph()
                add_inline_runs(paragraph, f"[缺少图片：{relative_path}]")
            i += 1
            continue

        if line.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows: list[list[str]] = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                rows.append(cells)
            add_table(document, rows)
            document.add_paragraph().paragraph_format.space_after = Pt(0)
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading_match:
            hashes, heading_text = heading_match.groups()
            level = len(hashes)
            if level == 1:
                if not first_h1_seen:
                    first_h1_seen = True
                else:
                    paragraph = document.add_paragraph(style="Heading 1")
                    add_inline_runs(paragraph, heading_text)
            else:
                style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
                paragraph = document.add_paragraph(style=style)
                add_inline_runs(paragraph, heading_text)
            i += 1
            continue

        if line.startswith(">"):
            quote_text = line.lstrip(">").strip()
            paragraph = document.add_paragraph(style="Quote")
            add_inline_runs(paragraph, quote_text)
            i += 1
            continue

        bullet_match = re.match(r"^[-*]\s+(.+)$", line)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet_match.group(1))
            i += 1
            continue

        numbered_match = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered_match:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, numbered_match.group(1))
            i += 1
            continue

        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, line)
        i += 1


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("Usage: build_docx.py <input.md> <output.docx>", file=sys.stderr)
        return 2

    input_path = Path(argv[1]).resolve()
    output_path = Path(argv[2]).resolve()
    if not input_path.is_file():
        print(f"Input not found: {input_path}", file=sys.stderr)
        return 1

    document = Document()
    configure_document(document)
    add_cover(document, "《逆岁：四契》")
    add_toc(document)
    parse_markdown(document, input_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
